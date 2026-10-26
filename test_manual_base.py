import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_manual():
    doc = Document()
    
    # ── Page Setup ────────────────────────────────────────────────────────────
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("CapFlow — Expense Tracker | Official Master User Manual (विस्तृत उपयोगकर्ता मार्गदर्शिका)")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("CapFlow • Private, Local-First Personal Finance & Splitter • https://capflow-rho.vercel.app/")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # ── Styling Helpers ───────────────────────────────────────────────────────
    PRIMARY_COLOR = RGBColor(30, 27, 75)      # Deep Indigo #1E1B4B
    ACCENT_COLOR = RGBColor(79, 70, 229)      # Brand Indigo #4F46E5
    SECONDARY_COLOR = RGBColor(14, 165, 233)  # Sky Blue #0EA5E9
    TEXT_DARK = RGBColor(15, 23, 42)          # Slate 900 #0F172A
    TEXT_MUTED = RGBColor(71, 85, 105)        # Slate 600 #475569
    HINDI_COLOR = RGBColor(51, 65, 85)        # Slate 700 #334155
    MOBILE_ACCENT = RGBColor(225, 29, 72)     # Rose 600 #E11D48

    def set_cell_shading(cell, color_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
        tblPr.append(borders)

    def add_callout(title, text, box_type="note"):
        colors = {
            "note": {"border": "4F46E5", "bg": "EEF2FF", "title_color": RGBColor(67, 56, 202)},
            "tip": {"border": "10B981", "bg": "ECFDF5", "title_color": RGBColor(6, 95, 70)},
            "warning": {"border": "EF4444", "bg": "FEF2F2", "title_color": RGBColor(153, 27, 27)},
            "example": {"border": "0EA5E9", "bg": "F0F9FF", "title_color": RGBColor(12, 74, 96)},
            "mobile": {"border": "E11D48", "bg": "FFF1F2", "title_color": RGBColor(159, 18, 57)}
        }
        cfg = colors.get(box_type, colors["note"])
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.9)
        set_cell_shading(cell, cfg["bg"])
        set_cell_margins(cell, top=110, bottom=110, left=170, right=150)
        
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{cfg["border"]}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        run_t = p.add_run(title + "\n")
        run_t.bold = True
        run_t.font.size = Pt(10)
        run_t.font.color.rgb = cfg["title_color"]
        run_t.font.name = "Arial"
        
        run_c = p.add_run(text)
        run_c.font.size = Pt(9.5)
        run_c.font.color.rgb = TEXT_DARK
        run_c.font.name = "Arial"
        
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def add_h1(text, hindi_text=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15.5)
        run.font.color.rgb = PRIMARY_COLOR
        run.font.name = "Arial"
        if hindi_text:
            run_hi = p.add_run(" | " + hindi_text)
            run_hi.bold = True
            run_hi.font.size = Pt(13.5)
            run_hi.font.color.rgb = ACCENT_COLOR
            run_hi.font.name = "Arial"

    def add_h2(text, hindi_text=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12.5)
        run.font.color.rgb = ACCENT_COLOR
        run.font.name = "Arial"
        if hindi_text:
            run_hi = p.add_run(" (" + hindi_text + ")")
            run_hi.font.size = Pt(11)
            run_hi.font.color.rgb = HINDI_COLOR
            run_hi.font.name = "Arial"

    def add_h3(text, hindi_text=None, is_mobile=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = MOBILE_ACCENT if is_mobile else PRIMARY_COLOR
        run.font.name = "Arial"
        if hindi_text:
            run_hi = p.add_run(" — " + hindi_text)
            run_hi.font.size = Pt(10)
            run_hi.font.color.rgb = HINDI_COLOR
            run_hi.font.name = "Arial"

    def add_p(en_text, hi_text=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(en_text)
        run.font.size = Pt(9.5)
        run.font.color.rgb = TEXT_DARK
        run.font.name = "Arial"
        if hi_text:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(5)
            p2.paragraph_format.line_spacing = 1.15
            run_hi = p2.add_run("हिंदी व्याख्या: " + hi_text)
            run_hi.font.size = Pt(9)
            run_hi.font.color.rgb = HINDI_COLOR
            run_hi.font.name = "Arial"

    def add_annotated_image(img_path, caption, width=Inches(6.2)):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(img_path, width=width)
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(6)
            run_cap = p_cap.add_run(caption)
            run_cap.font.size = Pt(8.5)
            run_cap.italic = True
            run_cap.font.color.rgb = TEXT_MUTED
            run_cap.font.name = "Arial"

    def add_callout_table(headers, rows_data, is_mobile=False):
        tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        
        hdr_color = "9F1239" if is_mobile else "1E1B4B"
        hdr_cells = tbl.rows[0].cells
        for idx, h_text in enumerate(headers):
            hdr_cells[idx].text = h_text
            set_cell_shading(hdr_cells[idx], hdr_color)
            set_cell_margins(hdr_cells[idx], top=70, bottom=70, left=90, right=90)
            p = hdr_cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.name = "Arial"
        
        for r_idx, row in enumerate(rows_data):
            row_cells = tbl.rows[r_idx + 1].cells
            bg_color = "FFF1F2" if (is_mobile and r_idx % 2 == 1) else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            for c_idx, val in enumerate(row):
                row_cells[c_idx].text = val
                set_cell_shading(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=90, right=90)
                p = row_cells[c_idx].paragraphs[0]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = TEXT_DARK
                    r.font.name = "Arial"
        
        tblPr = tbl._tbl.tblPr
        border_color = "FDA4AF" if is_mobile else "CBD5E1"
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/></w:tblBorders>')
        tblPr.append(borders)
        
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(5)

    print("Document initialized and helpers ready.")
    return doc, add_h1, add_h2, add_h3, add_p, add_annotated_image, add_callout, add_callout_table, set_cell_shading, set_cell_margins, set_table_borders

if __name__ == '__main__':
    create_manual()
