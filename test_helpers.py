import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def add_callout_box(doc, title, text, box_type="note"):
    colors = {
        "note": {"border": "4F46E5", "bg": "EEF2FF", "title_color": RGBColor(67, 56, 202)},
        "tip": {"border": "10B981", "bg": "ECFDF5", "title_color": RGBColor(6, 95, 70)},
        "warning": {"border": "EF4444", "bg": "FEF2F2", "title_color": RGBColor(153, 27, 27)},
        "info": {"border": "0EA5E9", "bg": "F0F9FF", "title_color": RGBColor(12, 74, 96)}
    }
    cfg = colors.get(box_type, colors["note"])
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, cfg["bg"])
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Left border only
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{cfg["border"]}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(title + "\n")
    run_t.bold = True
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = cfg["title_color"]
    run_t.font.name = "Arial"
    
    run_c = p.add_run(text)
    run_c.font.size = Pt(9.5)
    run_c.font.color.rgb = RGBColor(30, 41, 59)
    run_c.font.name = "Arial"
    
    # Space after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

print("Helper functions defined successfully")
