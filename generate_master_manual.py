# generate_master_manual.py
# Master Document Assembly Script for CapFlow
import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from manual_sections_core import build_core_chapters
from manual_sections_splitter import build_splitter_chapter
from manual_sections_features import build_features_chapters
from manual_sections_infra import build_infra_chapters

OUTPUT_DOCX = r"D:\webapp Project\Expense Tracker\app_capflow\CapFlow_Application_Manual.docx"
LOGO_PATH = r"nextjs/public/capflow-logo.png"

# Color Palette
PRIMARY_COLOR = RGBColor(30, 27, 75)      # Deep Indigo #1E1B4B
ACCENT_COLOR = RGBColor(79, 70, 229)      # Brand Indigo #4F46E5
SECONDARY_COLOR = RGBColor(14, 165, 233)  # Sky Blue #0EA5E9
TEXT_DARK = RGBColor(15, 23, 42)          # Slate 900 #0F172A
TEXT_MUTED = RGBColor(71, 85, 105)        # Slate 600 #475569
HINDI_COLOR = RGBColor(51, 65, 85)        # Slate 700 #334155
MOBILE_ACCENT = RGBColor(225, 29, 72)     # Rose 600 #E11D48

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def assemble_master_manual():
    print("Initializing CapFlow Master Manual Document Assembly...")
    doc = Document()
    
    # ── Page Setup ────────────────────────────────────────────────────────────
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        
        # Header
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("CapFlow — Expense Tracker | Official Master User Manual (विस्तृत उपयोगकर्ता मार्गदर्शिका)")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("CapFlow • Private, Local-First Personal Finance & Splitter • https://capflow-rho.vercel.app/")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # ── Typography & Layout Helpers ───────────────────────────────────────────
    def add_h1(text, hindi_text=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15.5)
        run.font.color.rgb = PRIMARY_COLOR
        run.font.name = "Arial"
        if hindi_text:
            run_hi = p.add_run(" | " + hindi_text)
            run_hi.bold = True
            run_hi.font.size = Pt(13)
            run_hi.font.color.rgb = ACCENT_COLOR
            run_hi.font.name = "Arial"

    def add_h2(text, hindi_text=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12.5)
        run.font.color.rgb = ACCENT_COLOR
        run.font.name = "Arial"
        if hindi_text:
            run_hi = p.add_run(" (" + hindi_text + ")")
            run_hi.font.size = Pt(11)
            run_hi.font.color.rgb = HINDI_COLOR
            run_hi.font.name = "Arial"

    def add_h3(text, hindi_text=None, is_mobile=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = MOBILE_ACCENT if is_mobile else PRIMARY_COLOR
        run.font.name = "Arial"
        if hindi_text:
            run_hi = p.add_run(" — " + hindi_text)
            run_hi.font.size = Pt(10)
            run_hi.font.color.rgb = HINDI_COLOR
            run_hi.font.name = "Arial"

    def add_p(en_text, hi_text=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(en_text)
        run.font.size = Pt(9.5)
        run.font.color.rgb = TEXT_DARK
        run.font.name = "Arial"
        if hi_text:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(5)
            p2.paragraph_format.line_spacing = 1.15
            run_hi = p2.add_run("हिंदी व्याख्या: " + hi_text)
            run_hi.font.size = Pt(9)
            run_hi.font.color.rgb = HINDI_COLOR
            run_hi.font.name = "Arial"

    def add_bullet(en_bold, en_rest, hi_text=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_b = p.add_run(en_bold + ": ")
        r_b.bold = True
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = PRIMARY_COLOR
        r_b.font.name = "Arial"
        r_r = p.add_run(en_rest)
        r_r.font.size = Pt(9.5)
        r_r.font.color.rgb = TEXT_DARK
        r_r.font.name = "Arial"
        if hi_text:
            r_hi = p.add_run(" — " + hi_text)
            r_hi.font.size = Pt(9)
            r_hi.font.color.rgb = HINDI_COLOR
            r_hi.font.name = "Arial"

    def add_callout(title, text, box_type="note"):
        colors = {
            "note": {"border": "4F46E5", "bg": "EEF2FF", "title_color": RGBColor(67, 56, 202)},
            "tip": {"border": "10B981", "bg": "ECFDF5", "title_color": RGBColor(6, 95, 70)},
            "warning": {"border": "EF4444", "bg": "FEF2F2", "title_color": RGBColor(153, 27, 27)},
            "example": {"border": "0EA5E9", "bg": "F0F9FF", "title_color": RGBColor(12, 74, 96)},
            "mobile": {"border": "E11D48", "bg": "FFF1F2", "title_color": RGBColor(159, 18, 57)}
        }
        cfg = colors.get(box_type, colors["note"])
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.9)
        set_cell_shading(cell, cfg["bg"])
        set_cell_margins(cell, top=110, bottom=110, left=170, right=150)
        
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{cfg["border"]}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        run_t = p.add_run(title + "\n")
        run_t.bold = True
        run_t.font.size = Pt(10)
        run_t.font.color.rgb = cfg["title_color"]
        run_t.font.name = "Arial"
        
        run_c = p.add_run(text)
        run_c.font.size = Pt(9.5)
        run_c.font.color.rgb = TEXT_DARK
        run_c.font.name = "Arial"
        
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def add_annotated_image(img_path, caption, width=Inches(6.2)):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(img_path, width=width)
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(6)
            run_cap = p_cap.add_run(caption)
            run_cap.font.size = Pt(8.5)
            run_cap.italic = True
            run_cap.font.color.rgb = TEXT_MUTED
            run_cap.font.name = "Arial"
        else:
            print(f"Warning: Image not found at {img_path}")

    def add_callout_table(headers, rows_data, is_mobile=False):
        tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        
        hdr_color = "9F1239" if is_mobile else "1E1B4B"
        hdr_cells = tbl.rows[0].cells
        for idx, h_text in enumerate(headers):
            hdr_cells[idx].text = h_text
            set_cell_shading(hdr_cells[idx], hdr_color)
            set_cell_margins(hdr_cells[idx], top=70, bottom=70, left=90, right=90)
            p = hdr_cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.name = "Arial"
        
        for r_idx, row in enumerate(rows_data):
            row_cells = tbl.rows[r_idx + 1].cells
            bg_color = "FFF1F2" if (is_mobile and r_idx % 2 == 1) else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            for c_idx, val in enumerate(row):
                row_cells[c_idx].text = val
                set_cell_shading(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=90, right=90)
                p = row_cells[c_idx].paragraphs[0]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = TEXT_DARK
                    r.font.name = "Arial"
        
        tblPr = tbl._tbl.tblPr
        border_color = "FDA4AF" if is_mobile else "CBD5E1"
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/></w:tblBorders>')
        tblPr.append(borders)
        
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(5)

    helpers = (add_h1, add_h2, add_h3, add_p, add_bullet, add_callout, add_annotated_image, add_callout_table)

    # =========================================================================
    # COVER PAGE
    # =========================================================================
    # Logo
    if os.path.exists(LOGO_PATH):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(18)
        p_logo.paragraph_format.space_after = Pt(8)
        r_logo = p_logo.add_run()
        r_logo.add_picture(LOGO_PATH, width=Inches(1.8))

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("CapFlow — Expense Tracker & Bill Splitter")
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = PRIMARY_COLOR
    r_title.font.name = "Arial"

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run("Official Comprehensive Master Application Manual & Reference Guide\n(विस्तृत उपयोगकर्ता मार्गदर्शिका एवं संपूर्ण संदर्भ ग्रंथ)")
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = ACCENT_COLOR
    r_sub.font.name = "Arial"

    # Metadata Table
    meta_headers = ["Attribute (विशेषता)", "Details & Operational Specifications (विवरण)"]
    meta_rows = [
        ["Application Title", "CapFlow (Personal Finance, Expense Tracker, EMI Manager & Group Splitter)"],
        ["Production URL", "https://capflow-rho.vercel.app/ (Zero-Config Progressive Web App)"],
        ["Software Release", "Version 1.0.0 Production Master (Build 2026.10-Universal)"],
        ["Privacy Architecture", "100% Client-Side Local-First (IndexedDB / Dexie.js) — Zero Server Data Storage"],
        ["Security Standard", "Hardware Biometrics (Touch ID / Face Unlock), 4-Digit Security PIN, Incognito Privacy Mask"],
        ["Splitter Capabilities", "Household/Family Mode, Headcount Weights, Dependent Routing, Session Budget Caps"],
        ["Settlement Engine", "Graph-Based Debt Simplification + Dynamic NPCI UPI QR & Intent Deep-Linking + Fullscreen Expander"],
        ["In-App Utility Tools", "Floating Calculator with History Tape, Mobile Camera Receipt Scanner, P2P Device Transfer"],
        ["Platform Support", "Cross-Platform: Windows, macOS, Linux, Android PWA, iOS Safari Home Screen"],
        ["Languages Included", "Fully Bilingual: English & हिन्दी (Detailed Step-by-Step Explanations)"],
        ["Total Visual Figures", "38 High-Resolution Figures (1 Official Logo + 19 Desktop Annotations + 18 Mobile Annotations)"]
    ]
    add_callout_table(meta_headers, meta_rows, is_mobile=False)

    # Key Feature Matrix Table
    add_h2("CapFlow Complete Feature Matrix & Architecture", "मुख्य फीचर मैट्रिक्स एवं तकनीकी अवलोकन")
    feature_headers = ["Subsystem / Feature", "Capability Scope", "Desktop Screen", "Mobile Screen"]
    feature_rows = [
        ["1. Authentication & Onboarding", "Google OAuth, Apple, Email OTP, Interactive Spotlight Tour", "Figure 1", "Figure 2"],
        ["2. Financial Dashboard", "Net Cash Flow Card, 7-Day Velocity Chart, Recent Feed, 7-Tab Dock", "Figure 3", "Figure 4"],
        ["3. Sub-3-Tap Expense Entry", "Amount Presets, 9 Categories, Mobile Camera Receipt Scanner", "Figure 5", "Figure 6"],
        ["4. Ledger & Statements", "Instant Sub-string Search, Type Filter Pills, Branded PDF/CSV Export", "Figure 7", "Figure 8"],
        ["5. Category Budgets & Goals", "Monthly Category Limits, 3-Tier Health Bars, Savings Goal Funds", "Figure 9", "Figure 10"],
        ["6. Interactive Calendar", "7x5 Monthly Heatmap Grid, Daily Cash Drill-Down Inspector", "Figure 11", "Figure 12"],
        ["7. Analytics & Intelligence", "Category Donut Charts, 6-Month Comparison Bars, Daily Trendline", "Figure 13", "Figure 14"],
        ["8. EMI & Debt Portfolio", "Loan Liability Overview, Tenure Progress (X/Y Paid), Due Reminders", "Figure 15", "Figure 16"],
        ["9. Bill Splitter Hub", "Trip & Flatmate Sessions, Join via 6-Char Code / Invite Link / QR", "Figure 19", "Figure 20"],
        ["10. Family & Household Mode", "Primary Payer Covers Dependents, Family Headcount Weights, Budget Caps", "Figure 33", "Figure 34"],
        ["11. 4 Split Distribution Modes", "Equal Split (=), Exact Amounts (₹), Percentages (%), Shares (x)", "Chapter 8.4", "Chapter 8.4"],
        ["12. Debt Simplification Engine", "Graph Greedy Reduction (Minimizes transactions to fewest net payments)", "Figure 21", "Figure 22"],
        ["13. Dynamic Settlement UPI QR", "NPCI UPI Deep-Link (GPay/PhonePe/Paytm/BHIM), Fullscreen QR Expander", "Figure 35", "Figure 36"],
        ["14. Floating In-App Calculator", "Real-Time Expression Parser, Parentheses, Tape, 1-Tap 'Apply Amount'", "Figure 31", "Figure 32"],
        ["15. Personal Payment QR Vault", "Upload & Auto-Decode UPI QR, Fullscreen Display, Instant VPA Copy", "Figure 29", "Figure 30"],
        ["16. Device-to-Device P2P Transfer", "Push Mode (15-min Token & QR) + Pull Mode (Live Camera QR Scanner)", "Figure 27", "Figure 28"],
        ["17. Appearance Console", "Light, Dark, OLED Midnight, 5 Accent Colors, Custom Hex, Font Scaling", "Figure 23", "Figure 24"],
        ["18. Privacy Center & App Lock", "Incognito Privacy Mask (••••), 4-Digit PIN, Biometrics, Storage Health", "Figure 25", "Figure 26"],
        ["19. Settings & Multi-Currency", "Indian Lakhs (₹1,00,000) vs Western Millions, Active Device Sessions", "Figure 17", "Figure 18"],
        ["20. Auto-Dispatch & Alerts", "Automated Monthly PDF Statement, 3-Day EMI Reminder, In-App Banners", "Figure 37", "Chapter 17"]
    ]
    add_callout_table(feature_headers, feature_rows, is_mobile=False)

    doc.add_page_break()

    # =========================================================================
    # EXECUTIVE SUMMARY & TABLE OF CONTENTS
    # =========================================================================
    add_h1("Executive Summary & Document Outline", "कार्यकारी सारांश एवं विषय सूची")
    add_p(
        "Welcome to the official master manual for CapFlow. This document serves as the exhaustive, authoritative guide "
        "to every single operational capability, mathematical algorithm, security protocol, and user interface workflow "
        "contained within the application. Every chapter is presented with bilingual English and Hindi descriptions, concrete examples, "
        "and dual-viewpoint annotated screenshots (both Desktop and Mobile) with red circle badges matching the on-screen UI components.",
        "CapFlow के आधिकारिक मास्टर मैनुअल में आपका स्वागत है। यह दस्तावेज़ ऐप की प्रत्येक कार्यप्रणाली, गणितीय एल्गोरिदम, सुरक्षा प्रणाली और यूजर इंटरफ़ेस का विस्तृत संदर्भ ग्रंथ है। प्रत्येक अध्याय अंग्रेजी और हिंदी दोनों भाषाओं में स्पष्ट उदाहरणों और डेस्कटॉप व मोबाइल दोनों के चिह्नित स्क्रीनशॉट्स (लाल गोलाकार बैज) के साथ प्रस्तुत किया गया है।"
    )

    toc_headers = ["Chapter #", "Chapter Title (English & Hindi)", "Key Focus Areas & Covered Features"]
    toc_rows = [
        ["Chapter 1", "Application Architecture & Privacy Foundation\n(आर्किटेक्चर एवं गोपनीयता आधार)", "Local-First IndexedDB, Dexie.js, Zero-Server Privacy, Sub-millisecond latency"],
        ["Chapter 2", "Getting Started & Secure Authentication\n(शुरुआत एवं सुरक्षित प्रमाणीकरण)", "Google OAuth, Apple, Facebook, Passwordless Email OTP, Figures 1 & 2"],
        ["Chapter 3", "Interactive Onboarding & Spotlight Tour\n(ऑनबोर्डिंग एवं लाइव स्पॉटलाइट टूर)", "WalkthroughProvider, 3-step feature discovery, on-demand tour replay"],
        ["Chapter 4", "Dashboard & Financial Command Center\n(डैशबोर्ड एवं वित्तीय नियंत्रण केंद्र)", "Cash Flow summary card, 7-day velocity chart, recent feed, 7-tab dock, Figures 3 & 4"],
        ["Chapter 5", "Logging Transactions & Camera Receipt Capture\n(लेन-देन प्रविष्टि एवं कैमरा रसीद)", "Sub-3-tap logging, amount presets, 9 visual categories, camera receipt capture, Figures 5 & 6"],
        ["Chapter 6", "Transaction Ledger, Search & Statement Export\n(लेन-देन इतिहास एवं स्टेटमेंट)", "Instant search, type filters, date drawer, bank-grade PDF & CSV statement exports, Figures 7 & 8"],
        ["Chapter 7", "Category Budgets & Savings Goals\n(बजट सीमाएं एवं बचत लक्ष्य)", "Category allowances, 3-tier health bars, target-date savings goals, Figures 9 & 10"],
        ["Chapter 8", "Bill Splitter & Group Expenses (Complete Subsystem)\n(बिल स्प्लिटर एवं समूह खर्च)", "Sessions hub, Family & Household Mode (Figs 33 & 34), 4 split modes, Floating Calculator (Figs 31 & 32), Debt Engine (Figs 21 & 22), Dynamic Settlement QR & Expander (Figs 35 & 36), Reports, Lock"],
        ["Chapter 9", "EMI & Loan Portfolio Tracking\n(ईएमआई एवं ऋण पोर्टफोलियो)", "Monthly debt burden, tenure progress (X/Y paid), automated 3-day advance alerts, Figures 15 & 16"],
        ["Chapter 10", "Interactive Financial Calendar & Daily Drill-Down\n(वित्तीय कैलेंडर एवं दैनिक विश्लेषण)", "7x5 monthly heatmap grid, day drill-down inspector, quick retrospective add, Figures 11 & 12"],
        ["Chapter 11", "Advanced Analytics & Spending Intelligence\n(उन्नत एनालिटिक्स एवं बुद्धिमत्ता)", "Category donut charts, 6-month historical comparison bars, daily trendline, Figures 13 & 14"],
        ["Chapter 12", "Personal Payment QR Code System\n(व्यक्तिगत भुगतान क्यूआर कोड)", "Upload personal UPI QR, client-side auto-decoding of VPA, fullscreen expand, Figures 29 & 30"],
        ["Chapter 13", "Device-to-Device Data Transfer & Account Migration\n(डिवाइस ट्रांसफर एवं माइग्रेशन)", "Push mode 15-min token/QR, Pull mode camera QR scanner, P2P database restore, Figures 27 & 28"],
        ["Chapter 14", "Appearance Customization Console\n(रूप-रंग एवं थीम अनुकूलन)", "Themes (Light/Dark/System/OLED Black), 5 Accents + Hex picker, 5 font scales, density, Figures 23 & 24"],
        ["Chapter 15", "Privacy Center, App Lock & Storage Health\n(गोपनीयता केंद्र एवं ऐप लॉक)", "Incognito privacy mask (••••), 4-digit PIN lock, WebAuthn biometrics, storage health, Figures 25 & 26"],
        ["Chapter 16", "Settings, Multi-Currency, Numbering & Sessions\n(सेटिंग्स एवं संख्या प्रणाली)", "Indian Lakhs (₹1,00,000) vs Western Millions, active device sessions tracker, Figures 17 & 18"],
        ["Chapter 17", "Web Push Notifications & Background Alerts\n(वेब पुश नोटिफिकेशन एवं बैकग्राउंड अलर्ट)", "Web Push API, background service workers, automated monthly PDF statement dispatch, 3-day EMI alerts, Figure 37"],
        ["Chapter 18", "Progressive Web App (PWA) & Offline Usability\n(प्रोग्रेसिव वेब ऐप एवं ऑफलाइन क्षमता)", "Android PWA install, iOS Safari Add to Home Screen, Desktop Chrome install, zero-network continuity"],
        ["Chapter 19", "About & Legal Support Hub\n(हमारे बारे में एवं सहायता केंद्र)", "Help Center, Privacy Policy, Terms of Service, Version 1.0.0 release notes"],
        ["Chapter 20", "Comprehensive FAQ & Troubleshooting Guide\n(सामान्य प्रश्न एवं समस्या निवारण)", "12 detailed bilingual Q&As covering privacy, storage resets, P2P Transfers, Split math, UPI links, Exports"]
    ]
    add_callout_table(toc_headers, toc_rows, is_mobile=False)

    doc.add_page_break()

    # =========================================================================
    # BUILD CHAPTERS
    # =========================================================================
    print("Building Chapters 1 to 7 (Core Ledger & Budgets)...")
    build_core_chapters(doc, helpers)

    print("Building Chapter 8 (Complete Bill Splitter & Floating Calculator)...")
    build_splitter_chapter(doc, helpers)

    print("Building Chapters 9 to 16 (Features, EMIs, Calendar, QR, Transfer, Privacy)...")
    build_features_chapters(doc, helpers)

    print("Building Chapters 17 to 20 (Notifications, PWA, Legal, FAQs)...")
    build_infra_chapters(doc, helpers)

    # ── Final Save ────────────────────────────────────────────────────────────
    print(f"Saving finalized master manual to: {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    file_size = os.path.getsize(OUTPUT_DOCX)
    print(f"SUCCESS: CapFlow Master User Manual generated successfully! File size: {file_size / (1024*1024):.2f} MB")

if __name__ == '__main__':
    assemble_master_manual()
